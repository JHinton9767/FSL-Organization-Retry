from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Sequence

import pandas as pd
from openpyxl import Workbook

from src.build_executive_report import clean_text, coerce_numeric, load_latest_bundle, yes_mask
from src.build_master_roster import autosize_columns, style_header


ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SNAPSHOT_ROOT = ROOT / "data" / "inbox" / "academic"
DEFAULT_OUTPUT_ROOT = ROOT / "output" / "current_snapshot_metrics"
UNRESOLVED_OUTCOMES = {"Active/Unknown", "No Further Observation", "Unknown", ""}
SNAPSHOT_REQUIRED_COLUMNS = {"Student ID", "First Name", "Last Name"}

SNAPSHOT_ALIAS_GROUPS = {
    "Student ID": {"student id", "banner id", "banner", "student number"},
    "First Name": {"first name", "firstname"},
    "Last Name": {"last name", "lastname"},
    "NetID": {"netid", "net id"},
    "High School GPA": {"high school gpa", "hs gpa"},
    "Overall GPA": {"overall gpa"},
    "Institutional GPA": {"institutional gpa", "txst gpa", "texas state gpa"},
    "Transfer GPA": {"transfer gpa"},
    "Total Credit Hours": {"total credit hours", "total hours", "overall credit hours"},
    "TXST Credit Hours": {"txst credit hours", "texas state credit hours", "institutional credit hours"},
    "Previous Semester GPA": {"previous semester gpa", "previous term gpa"},
    "Student Status": {"student status"},
    "Student Status (FT/PT)": {"student status ftpt", "student status ft pt", "student status (ft/pt)", "ft/pt"},
}

SNAPSHOT_COLUMNS = [
    "Student ID",
    "First Name",
    "Last Name",
    "NetID",
    "High School GPA",
    "Overall GPA",
    "Institutional GPA",
    "Transfer GPA",
    "Total Credit Hours",
    "TXST Credit Hours",
    "Previous Semester GPA",
    "Student Status",
    "Student Status (FT/PT)",
]

AUGMENTED_SUMMARY_COLUMNS = [
    "Student ID",
    "Preferred Last Name",
    "Preferred First Name",
    "Initial Chapter",
    "Organization Entry Cohort",
    "First Observed Organization Term",
    "Last Observed Organization Term",
    "First Observed Academic Term",
    "Latest Known Outcome Bucket",
    "Snapshot Matched",
    "Snapshot Student Status",
    "Snapshot Student Status (FT/PT)",
    "Snapshot High School GPA",
    "Snapshot Overall GPA",
    "Snapshot Institutional GPA",
    "Snapshot Transfer GPA",
    "Snapshot Total Credit Hours",
    "Snapshot TXST Credit Hours",
    "Snapshot Previous Semester GPA",
    "Estimated Transfer Credit Hours",
    "Observed Passed Hours Since Org Entry",
    "Estimated Pre-Org Credit Hours (Overall Basis)",
    "Estimated Pre-Org Credit Hours (TXST Basis)",
    "Estimated Pre-Org Stage (Overall Basis)",
    "Estimated Pre-Org Stage (TXST Basis)",
    "Augmented Latest Outcome Bucket",
    "Augmented Ever Graduated Flag",
    "Current Enrollment Intensity Bucket",
]

COHORT_METRIC_COLUMNS = [
    "Metric Group",
    "Metric Label",
    "Cohort",
    "Eligible Students",
    "Student Count",
    "Rate",
    "Average Value",
    "Notes",
]

CHAPTER_METRIC_COLUMNS = [
    "Chapter",
    "Students",
    "Snapshot Matches",
    "Snapshot Match Rate",
    "Augmented Graduation Rate (Resolved)",
    "Average Snapshot Institutional GPA",
    "Average Snapshot Overall GPA",
    "Average Snapshot Total Credit Hours",
    "Average Snapshot TXST Credit Hours",
]

QA_COLUMNS = ["Check", "Value", "Notes"]
DEFINITION_COLUMNS = ["Field", "Definition"]


@dataclass(frozen=True)
class SnapshotAnalyticsResult:
    output_folder: Path
    workbook_path: Path
    augmented_summary_csv: Path
    cohort_metrics_csv: Path
    chapter_metrics_csv: Path
    qa_csv: Path
    methodology_path: Path
    snapshot_sources: List[Path]
    enhanced_folder: Path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Additively merge a current one-row-per-student academic snapshot into the existing enhanced analytics bundle "
            "to improve outcome resolution and pre-organization timing estimates."
        )
    )
    parser.add_argument("--enhanced-root", default=str(ROOT / "output" / "enhanced_metrics"))
    parser.add_argument("--enhanced-folder", default="")
    parser.add_argument("--enhanced-workbook", default="")
    parser.add_argument("--snapshot-root", default=str(DEFAULT_SNAPSHOT_ROOT))
    parser.add_argument("--snapshot", default="")
    parser.add_argument("--output-root", default=str(DEFAULT_OUTPUT_ROOT))
    return parser.parse_args()


def canonical_header(value: object) -> str:
    text = clean_text(value).lower().replace("_", " ")
    text = "".join(char if char.isalnum() or char == " " else " " for char in text)
    return " ".join(text.split())


def normalize_student_id(value: object) -> str:
    text = clean_text(value)
    if text.endswith(".0"):
        text = text[:-2]
    return text


def is_snapshot_filename(path: Path) -> bool:
    stem = clean_text(path.stem).lower()
    return stem.startswith("new member")


def infer_snapshot_paths(explicit_snapshot: Optional[Path], snapshot_root: Path) -> List[Path]:
    if explicit_snapshot:
        snapshot_path = explicit_snapshot.expanduser().resolve()
        if not snapshot_path.exists():
            raise FileNotFoundError(f"Snapshot file not found: {snapshot_path}")
        return [snapshot_path]
    if not snapshot_root.exists():
        raise FileNotFoundError(
            f"Snapshot folder not found at {snapshot_root}. Place the 'New Member (X)' files there or pass --snapshot."
        )
    candidates = sorted(
        [
            path
            for path in snapshot_root.iterdir()
            if path.is_file()
            and path.suffix.lower() in {'.xlsx', '.xlsm', '.csv'}
            and is_snapshot_filename(path)
        ],
        key=lambda item: (item.stat().st_mtime, item.name.lower()),
    )
    if not candidates:
        raise FileNotFoundError(
            f"No 'New Member' snapshot files found in {snapshot_root}. Place files named like 'New Member (1).xlsx' there or pass --snapshot."
        )
    return candidates


def rename_snapshot_columns(columns: Sequence[object]) -> List[str]:
    renamed: List[str] = []
    for column in columns:
        canonical = canonical_header(column)
        match = next(
            (
                standard
                for standard, aliases in SNAPSHOT_ALIAS_GROUPS.items()
                if canonical in aliases
            ),
            clean_text(column),
        )
        renamed.append(match)
    return renamed


def choose_best_snapshot_rows(frame: pd.DataFrame) -> pd.DataFrame:
    if frame.empty:
        return frame
    frame = frame.copy()
    frame["_filled_fields"] = frame.notna().sum(axis=1)
    frame["_status_present"] = frame["Student Status"].fillna("").astype(str).str.strip().ne("").astype(int)
    frame = frame.sort_values(
        by=["Student ID", "_filled_fields", "_status_present", "Last Name", "First Name"],
        ascending=[True, False, False, True, True],
    )
    frame = frame.drop_duplicates(subset=["Student ID"], keep="first")
    return frame.drop(columns=["_filled_fields", "_status_present"])


def load_snapshot_table(snapshot_path: Path) -> pd.DataFrame:
    if snapshot_path.suffix.lower() == ".csv":
        raw = pd.read_csv(snapshot_path)
        raw.columns = rename_snapshot_columns(raw.columns)
        frame = raw
    else:
        sheet_map = pd.read_excel(snapshot_path, sheet_name=None)
        selected: Optional[pd.DataFrame] = None
        for _, frame in sheet_map.items():
            candidate = frame.copy()
            candidate.columns = rename_snapshot_columns(candidate.columns)
            if SNAPSHOT_REQUIRED_COLUMNS.issubset(set(candidate.columns)):
                selected = candidate
                break
        if selected is None:
            raise FileNotFoundError(
                f"No usable sheet with required columns ({', '.join(sorted(SNAPSHOT_REQUIRED_COLUMNS))}) was found in {snapshot_path}."
            )
        frame = selected

    for column in SNAPSHOT_COLUMNS:
        if column not in frame.columns:
            frame[column] = ""
    frame = frame[SNAPSHOT_COLUMNS].copy()
    for column in ["Student ID", "First Name", "Last Name", "NetID", "Student Status", "Student Status (FT/PT)"]:
        frame[column] = frame[column].map(clean_text)
    frame["Student ID"] = frame["Student ID"].map(normalize_student_id)
    frame = frame.loc[frame["Student ID"].ne("")]
    frame = choose_best_snapshot_rows(frame)
    return frame.reset_index(drop=True)


def load_combined_snapshot_table(snapshot_paths: Sequence[Path]) -> pd.DataFrame:
    combined_frames: List[pd.DataFrame] = []
    for snapshot_path in snapshot_paths:
        frame = load_snapshot_table(snapshot_path)
        frame["Snapshot Source File"] = snapshot_path.name
        combined_frames.append(frame)
    if not combined_frames:
        return pd.DataFrame(columns=SNAPSHOT_COLUMNS + ["Snapshot Source File"])
    combined = pd.concat(combined_frames, ignore_index=True)
    combined = choose_best_snapshot_rows(combined)
    return combined.reset_index(drop=True)


def bucket_30_hours(value: object) -> str:
    if value == "" or pd.isna(value):
        return "Unknown"
    number = float(value)
    if number < 0:
        return "Unknown"
    lower = int(number // 30) * 30
    upper = lower + 29
    return f"{lower}-{upper}"


def current_intensity_bucket(value: object) -> str:
    text = canonical_header(value)
    if not text:
        return "Unknown"
    if "full" in text or text == "ft":
        return "Full-Time"
    if "part" in text or text == "pt":
        return "Part-Time"
    return clean_text(value)


def snapshot_outcome_bucket(status_value: object) -> str:
    text = canonical_header(status_value)
    if not text:
        return "Unknown"
    if any(word in text for word in ["graduat", "degree", "complete", "completed", "alumni"]):
        return "Graduated"
    if "suspend" in text:
        return "Suspended"
    if "transfer" in text:
        return "Transfer"
    if any(word in text for word in ["inactive", "drop", "withdraw", "dismiss", "separat", "revoke", "resign"]):
        return "Dropped/Resigned/Revoked/Inactive"
    if any(word in text for word in ["active", "current", "enrolled", "good standing", "probation", "warning", "ft", "pt"]):
        return "Active/Enrolled"
    return "Unknown"


def build_observed_hours_lookup(longitudinal: pd.DataFrame) -> pd.DataFrame:
    frame = longitudinal.copy()
    frame["Student ID"] = frame["Student ID"].map(normalize_student_id)
    passed_hours = coerce_numeric(frame["Term Passed Hours"]) if "Term Passed Hours" in frame.columns else pd.Series(index=frame.index, dtype=float)
    relative_term = coerce_numeric(frame["Relative Term Index From Org Entry"]) if "Relative Term Index From Org Entry" in frame.columns else pd.Series(index=frame.index, dtype=float)
    academic_present = yes_mask(frame["Academic Present"]) if "Academic Present" in frame.columns else pd.Series(False, index=frame.index)
    eligible = frame.loc[academic_present & relative_term.ge(0, fill_value=False)].copy()
    eligible["Term Passed Hours"] = passed_hours.loc[eligible.index].fillna(0.0)
    grouped = eligible.groupby("Student ID", dropna=False)["Term Passed Hours"].sum().reset_index()
    grouped = grouped.rename(columns={"Term Passed Hours": "Observed Passed Hours Since Org Entry"})
    return grouped


def merge_augmented_summary(summary: pd.DataFrame, longitudinal: pd.DataFrame, snapshot: pd.DataFrame) -> pd.DataFrame:
    frame = summary.copy()
    frame["Student ID"] = frame["Student ID"].map(normalize_student_id)
    observed_hours = build_observed_hours_lookup(longitudinal)
    merged = frame.merge(snapshot, on="Student ID", how="left", suffixes=("", " Snapshot"))
    merged = merged.merge(observed_hours, on="Student ID", how="left")
    merged["Observed Passed Hours Since Org Entry"] = coerce_numeric(merged["Observed Passed Hours Since Org Entry"]).fillna(0.0)

    merged["Snapshot Matched"] = merged["First Name"].fillna("").astype(str).str.strip().ne("").map(lambda value: "Yes" if value else "No")
    rename_map = {
        "First Name": "Snapshot First Name",
        "Last Name": "Snapshot Last Name",
        "NetID": "Snapshot NetID",
        "High School GPA": "Snapshot High School GPA",
        "Overall GPA": "Snapshot Overall GPA",
        "Institutional GPA": "Snapshot Institutional GPA",
        "Transfer GPA": "Snapshot Transfer GPA",
        "Total Credit Hours": "Snapshot Total Credit Hours",
        "TXST Credit Hours": "Snapshot TXST Credit Hours",
        "Previous Semester GPA": "Snapshot Previous Semester GPA",
        "Student Status": "Snapshot Student Status",
        "Student Status (FT/PT)": "Snapshot Student Status (FT/PT)",
    }
    merged = merged.rename(columns=rename_map)

    for column in [
        "Snapshot High School GPA",
        "Snapshot Overall GPA",
        "Snapshot Institutional GPA",
        "Snapshot Transfer GPA",
        "Snapshot Total Credit Hours",
        "Snapshot TXST Credit Hours",
        "Snapshot Previous Semester GPA",
    ]:
        merged[column] = coerce_numeric(merged[column])

    merged["Estimated Transfer Credit Hours"] = (
        merged["Snapshot Total Credit Hours"].fillna(0.0) - merged["Snapshot TXST Credit Hours"].fillna(0.0)
    ).clip(lower=0.0)
    merged["Estimated Pre-Org Credit Hours (Overall Basis)"] = (
        merged["Snapshot Total Credit Hours"].fillna(0.0) - merged["Observed Passed Hours Since Org Entry"].fillna(0.0)
    ).clip(lower=0.0)
    merged["Estimated Pre-Org Credit Hours (TXST Basis)"] = (
        merged["Snapshot TXST Credit Hours"].fillna(0.0) - merged["Observed Passed Hours Since Org Entry"].fillna(0.0)
    ).clip(lower=0.0)
    merged["Estimated Pre-Org Stage (Overall Basis)"] = merged["Estimated Pre-Org Credit Hours (Overall Basis)"].map(bucket_30_hours)
    merged["Estimated Pre-Org Stage (TXST Basis)"] = merged["Estimated Pre-Org Credit Hours (TXST Basis)"].map(bucket_30_hours)
    merged["Current Enrollment Intensity Bucket"] = merged["Snapshot Student Status (FT/PT)"].map(current_intensity_bucket)
    merged["Snapshot Explicit Outcome Bucket"] = merged["Snapshot Student Status"].map(snapshot_outcome_bucket)

    def augmented_outcome(row: pd.Series) -> str:
        existing = clean_text(row.get("Latest Known Outcome Bucket"))
        if existing not in UNRESOLVED_OUTCOMES:
            return existing
        snapshot_bucket = clean_text(row.get("Snapshot Explicit Outcome Bucket"))
        if snapshot_bucket in {"Graduated", "Suspended", "Transfer", "Dropped/Resigned/Revoked/Inactive"}:
            return snapshot_bucket
        if snapshot_bucket == "Active/Enrolled":
            return "Active/Unknown"
        return existing or "Unknown"

    merged["Augmented Latest Outcome Bucket"] = merged.apply(augmented_outcome, axis=1)
    merged["Augmented Ever Graduated Flag"] = merged.apply(
        lambda row: "Yes"
        if clean_text(row.get("Ever Graduated Flag")) == "Yes" or clean_text(row.get("Augmented Latest Outcome Bucket")) == "Graduated"
        else "No",
        axis=1,
    )
    return merged


def mean_or_blank(series: pd.Series) -> object:
    numeric = coerce_numeric(series)
    if numeric.dropna().empty:
        return ""
    return float(numeric.dropna().mean())


def add_metric_row(
    rows: List[Dict[str, object]],
    metric_group: str,
    metric_label: str,
    cohort: str,
    eligible_students: object = "",
    student_count: object = "",
    rate: object = "",
    average_value: object = "",
    notes: str = "",
) -> None:
    rows.append(
        {
            "Metric Group": metric_group,
            "Metric Label": metric_label,
            "Cohort": cohort,
            "Eligible Students": eligible_students,
            "Student Count": student_count,
            "Rate": rate,
            "Average Value": average_value,
            "Notes": notes,
        }
    )


def build_augmented_cohort_metrics(summary: pd.DataFrame) -> pd.DataFrame:
    rows: List[Dict[str, object]] = []
    groups: Dict[str, pd.DataFrame] = {"Overall": summary}
    for cohort in sorted({clean_text(value) for value in summary["Organization Entry Cohort"].tolist() if clean_text(value)}):
        groups[cohort] = summary.loc[summary["Organization Entry Cohort"].fillna("").astype(str).str.strip().eq(cohort)].copy()

    for cohort, frame in groups.items():
        if frame.empty:
            continue
        add_metric_row(rows, "Coverage", "Students", cohort, student_count=len(frame))
        snapshot_matches = yes_mask(frame["Snapshot Matched"]).sum()
        add_metric_row(
            rows,
            "Coverage",
            "Snapshot match rate",
            cohort,
            eligible_students=len(frame),
            student_count=int(snapshot_matches),
            rate=float(snapshot_matches) / float(len(frame)),
            notes="Share of students matched to the current one-row-per-student academic snapshot.",
        )

        resolved = frame.loc[
            frame["Augmented Latest Outcome Bucket"].fillna("").astype(str).str.strip().isin(
                ["Graduated", "Suspended", "Transfer", "Dropped/Resigned/Revoked/Inactive"]
            )
        ]
        grad_count = int(resolved["Augmented Latest Outcome Bucket"].fillna("").astype(str).str.strip().eq("Graduated").sum())
        add_metric_row(
            rows,
            "Outcomes",
            "Augmented eventual graduation rate from organization entry",
            cohort,
            eligible_students=len(resolved),
            student_count=grad_count,
            rate=(float(grad_count) / float(len(resolved))) if len(resolved) else "",
            notes="Uses the existing observed outcome when resolved, otherwise a cautiously parsed current snapshot status when that status is explicit.",
        )

        for label, field in [
            ("Average current institutional GPA", "Snapshot Institutional GPA"),
            ("Average current overall GPA", "Snapshot Overall GPA"),
            ("Average current previous semester GPA", "Snapshot Previous Semester GPA"),
            ("Average current total credit hours", "Snapshot Total Credit Hours"),
            ("Average current TXST credit hours", "Snapshot TXST Credit Hours"),
            ("Average estimated pre-org credit hours (TXST basis)", "Estimated Pre-Org Credit Hours (TXST Basis)"),
        ]:
            add_metric_row(
                rows,
                "Current Snapshot Academics",
                label,
                cohort,
                eligible_students=int(coerce_numeric(frame[field]).notna().sum()),
                average_value=mean_or_blank(frame[field]),
            )

        for intensity in ["Full-Time", "Part-Time"]:
            count = int(frame["Current Enrollment Intensity Bucket"].fillna("").astype(str).str.strip().eq(intensity).sum())
            add_metric_row(
                rows,
                "Current Snapshot Academics",
                f"Current {intensity.lower()} share",
                cohort,
                eligible_students=len(frame),
                student_count=count,
                rate=(float(count) / float(len(frame))) if len(frame) else "",
            )
    return pd.DataFrame(rows, columns=COHORT_METRIC_COLUMNS)


def build_chapter_metrics(summary: pd.DataFrame) -> pd.DataFrame:
    rows: List[List[object]] = []
    for chapter in sorted({clean_text(value) for value in summary["Initial Chapter"].tolist() if clean_text(value)}):
        frame = summary.loc[summary["Initial Chapter"].fillna("").astype(str).str.strip().eq(chapter)].copy()
        if frame.empty:
            continue
        snapshot_matches = int(yes_mask(frame["Snapshot Matched"]).sum())
        resolved = frame.loc[
            frame["Augmented Latest Outcome Bucket"].fillna("").astype(str).str.strip().isin(
                ["Graduated", "Suspended", "Transfer", "Dropped/Resigned/Revoked/Inactive"]
            )
        ]
        grad_rate = ""
        if not resolved.empty:
            grad_rate = float(
                resolved["Augmented Latest Outcome Bucket"].fillna("").astype(str).str.strip().eq("Graduated").sum()
            ) / float(len(resolved))
        rows.append(
            [
                chapter,
                len(frame),
                snapshot_matches,
                (float(snapshot_matches) / float(len(frame))) if len(frame) else "",
                grad_rate,
                mean_or_blank(frame["Snapshot Institutional GPA"]),
                mean_or_blank(frame["Snapshot Overall GPA"]),
                mean_or_blank(frame["Snapshot Total Credit Hours"]),
                mean_or_blank(frame["Snapshot TXST Credit Hours"]),
            ]
        )
    return pd.DataFrame(rows, columns=CHAPTER_METRIC_COLUMNS)


def build_qa_table(summary: pd.DataFrame, snapshot: pd.DataFrame) -> pd.DataFrame:
    duplicate_snapshot_ids = int(snapshot["Student ID"].fillna("").astype(str).str.strip().duplicated().sum())
    matched = int(yes_mask(summary["Snapshot Matched"]).sum())
    unresolved_before = int(summary["Latest Known Outcome Bucket"].fillna("").astype(str).str.strip().isin(UNRESOLVED_OUTCOMES).sum())
    unresolved_after = int(summary["Augmented Latest Outcome Bucket"].fillna("").astype(str).str.strip().isin(UNRESOLVED_OUTCOMES).sum())
    txst_gt_total = int(
        (
            coerce_numeric(summary["Snapshot TXST Credit Hours"]).fillna(-1)
            > coerce_numeric(summary["Snapshot Total Credit Hours"]).fillna(-1)
        ).sum()
    )
    return pd.DataFrame(
        [
            ["Snapshot rows", len(snapshot), "Rows kept after Student ID dedupe."],
            ["Duplicate snapshot Student IDs removed", duplicate_snapshot_ids, "Should ideally be 0 after preprocessing."],
            ["Enhanced students", len(summary), "Rows in the existing student summary."],
            ["Snapshot matches", matched, "Students matched by Student ID."],
            [
                "Unresolved outcomes before snapshot augmentation",
                unresolved_before,
                "Count of Active/Unknown, No Further Observation, Unknown, or blank outcomes in the original summary.",
            ],
            [
                "Unresolved outcomes after snapshot augmentation",
                unresolved_after,
                "Count after cautiously applying explicit current snapshot outcomes when available.",
            ],
            ["Rows where TXST credit hours exceeded total credit hours", txst_gt_total, "Should normally be 0."],
        ],
        columns=QA_COLUMNS,
    )


def build_definitions_table() -> pd.DataFrame:
    return pd.DataFrame(
        [
            ["Snapshot match", "Student matched from the enhanced student summary to the current one-row-per-student academic snapshot using Student ID."],
            ["Augmented latest outcome bucket", "The original latest observed outcome unless it was unresolved, in which case an explicit current snapshot status may resolve it."],
            ["Estimated pre-org credit hours (TXST basis)", "Current TXST credit hours minus observed passed hours during the organization-observation window. This is an estimate, not a historical ledger."],
            ["Estimated pre-org stage", "A 30-credit-hour bucket used to approximate how far along a student was before observed organization entry."],
            ["Augmented eventual graduation rate from organization entry", "Graduation rate using the original resolved outcome when available, plus explicit current snapshot graduation statuses when they can be parsed safely."],
        ],
        columns=DEFINITION_COLUMNS,
    )


def write_records_sheet(wb: Workbook, sheet_name: str, frame: pd.DataFrame) -> None:
    ws = wb.create_sheet(title=sheet_name[:31])
    ws.append(list(frame.columns))
    style_header(ws)
    for row in frame.itertuples(index=False, name=None):
        ws.append(list(row))
    ws.freeze_panes = "A2"
    autosize_columns(ws)


def write_overview_sheet(
    wb: Workbook,
    summary: pd.DataFrame,
    snapshot_sources: Sequence[Path],
    enhanced_folder: Path,
) -> None:
    ws = wb.active
    ws.title = "Overview"
    ws.append(["Metric", "Value"])
    style_header(ws)

    matched = int(yes_mask(summary["Snapshot Matched"]).sum())
    resolved = summary.loc[
        summary["Augmented Latest Outcome Bucket"].fillna("").astype(str).str.strip().isin(
            ["Graduated", "Suspended", "Transfer", "Dropped/Resigned/Revoked/Inactive"]
        )
    ]
    grad_rate = ""
    if not resolved.empty:
        grad_rate = float(
            resolved["Augmented Latest Outcome Bucket"].fillna("").astype(str).str.strip().eq("Graduated").sum()
        ) / float(len(resolved))
    snapshot_label = (
        str(snapshot_sources[0]) if len(snapshot_sources) == 1 else f"{len(snapshot_sources)} files from {snapshot_sources[0].parent}"
    )
    metrics = [
        ["Enhanced analytics source folder", str(enhanced_folder)],
        ["Snapshot source files", snapshot_label],
        ["Students in enhanced summary", len(summary)],
        ["Snapshot matches", matched],
        ["Snapshot match rate", (float(matched) / float(len(summary))) if len(summary) else ""],
        ["Augmented graduation rate from organization entry", grad_rate],
        ["Average current institutional GPA", mean_or_blank(summary["Snapshot Institutional GPA"])],
        ["Average current overall GPA", mean_or_blank(summary["Snapshot Overall GPA"])],
        ["Average current total credit hours", mean_or_blank(summary["Snapshot Total Credit Hours"])],
        ["Average estimated pre-org credit hours (TXST basis)", mean_or_blank(summary["Estimated Pre-Org Credit Hours (TXST Basis)"])],
    ]
    for row in metrics:
        ws.append(row)
    ws.freeze_panes = "A2"
    autosize_columns(ws)


def write_methodology(path: Path, snapshot_sources: Sequence[Path], enhanced_folder: Path) -> None:
    if len(snapshot_sources) == 1:
        snapshot_line = f"- Current snapshot source: `{snapshot_sources[0]}`"
    else:
        snapshot_line = f"- Current snapshot sources: `{len(snapshot_sources)}` files from `{snapshot_sources[0].parent}`"
    lines = [
        "# Current Snapshot Augmentation Methodology",
        "",
        f"- Enhanced analytics source: `{enhanced_folder}`",
        snapshot_line,
        "",
        "## What this adds",
        "",
        "- Combines every current snapshot file named like `New Member (X)` that was found in the selected folder.",
        "- Merges a one-row-per-student current academic snapshot onto the existing enhanced student summary by Student ID.",
        "- Uses explicit current snapshot statuses, when they can be parsed safely, to resolve some previously unresolved outcomes.",
        "- Estimates pre-organization credit accumulation by subtracting observed passed hours during the organization window from current total or TXST credit hours.",
        "",
        "## What this does not claim",
        "",
        "- The current snapshot is not historical term-by-term data.",
        "- Pre-organization credit estimates are approximations, not exact reconstructed academic histories.",
        "- Exact school-entry terms still require explicit enrollment history or registrar-level entry dates.",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")


def build_current_snapshot_analytics(
    enhanced_root: Path,
    explicit_enhanced_folder: Optional[Path],
    explicit_enhanced_workbook: Optional[Path],
    snapshot_root: Path,
    explicit_snapshot: Optional[Path],
    output_root: Path,
) -> SnapshotAnalyticsResult:
    bundle = load_latest_bundle(
        enhanced_root=enhanced_root,
        explicit_folder=explicit_enhanced_folder,
        explicit_workbook=explicit_enhanced_workbook,
    )
    if "master_longitudinal" not in bundle.tables:
        raise FileNotFoundError(
            "The enhanced analytics bundle does not include Master_Longitudinal. Run py run_enhanced_org_analytics.py first."
        )
    snapshot_sources = infer_snapshot_paths(explicit_snapshot, snapshot_root)
    snapshot = load_combined_snapshot_table(snapshot_sources)
    summary = bundle.tables["student_summary"].copy()
    longitudinal = bundle.tables["master_longitudinal"].copy()

    augmented_summary = merge_augmented_summary(summary, longitudinal, snapshot)
    cohort_metrics = build_augmented_cohort_metrics(augmented_summary)
    chapter_metrics = build_chapter_metrics(augmented_summary)
    qa_table = build_qa_table(augmented_summary, snapshot)
    definitions = build_definitions_table()

    timestamp = datetime.now().strftime("run_%Y%m%d_%H%M%S")
    output_folder = output_root / timestamp
    output_folder.mkdir(parents=True, exist_ok=True)

    workbook_path = output_folder / f"organization_entry_snapshot_augmented_{timestamp}.xlsx"
    augmented_summary_csv = output_folder / "snapshot_augmented_student_summary.csv"
    cohort_metrics_csv = output_folder / "snapshot_augmented_cohort_metrics.csv"
    chapter_metrics_csv = output_folder / "snapshot_augmented_chapter_metrics.csv"
    qa_csv = output_folder / "snapshot_merge_qa.csv"
    methodology_path = output_folder / "methodology.md"

    augmented_summary[AUGMENTED_SUMMARY_COLUMNS].to_csv(augmented_summary_csv, index=False)
    cohort_metrics.to_csv(cohort_metrics_csv, index=False)
    chapter_metrics.to_csv(chapter_metrics_csv, index=False)
    qa_table.to_csv(qa_csv, index=False)
    write_methodology(methodology_path, snapshot_sources, bundle.enhanced_folder)

    wb = Workbook()
    write_overview_sheet(wb, augmented_summary, snapshot_sources, bundle.enhanced_folder)
    write_records_sheet(wb, "Augmented_Summary", augmented_summary[AUGMENTED_SUMMARY_COLUMNS])
    write_records_sheet(wb, "Cohort_Metrics", cohort_metrics)
    write_records_sheet(wb, "Chapter_Metrics", chapter_metrics)
    write_records_sheet(
        wb,
        "Join_Timing",
        augmented_summary[
            [
                "Student ID",
                "Preferred Last Name",
                "Preferred First Name",
                "Initial Chapter",
                "Organization Entry Cohort",
                "Entry Cumulative Hours",
                "Snapshot Total Credit Hours",
                "Snapshot TXST Credit Hours",
                "Observed Passed Hours Since Org Entry",
                "Estimated Pre-Org Credit Hours (Overall Basis)",
                "Estimated Pre-Org Credit Hours (TXST Basis)",
                "Estimated Pre-Org Stage (Overall Basis)",
                "Estimated Pre-Org Stage (TXST Basis)",
                "Augmented Latest Outcome Bucket",
            ]
        ],
    )
    write_records_sheet(wb, "Snapshot_QA", qa_table)
    write_records_sheet(wb, "Definitions", definitions)
    write_records_sheet(wb, "Unmatched_Snapshot", snapshot.loc[~snapshot["Student ID"].isin(set(augmented_summary["Student ID"]))])
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(workbook_path)

    return SnapshotAnalyticsResult(
        output_folder=output_folder,
        workbook_path=workbook_path,
        augmented_summary_csv=augmented_summary_csv,
        cohort_metrics_csv=cohort_metrics_csv,
        chapter_metrics_csv=chapter_metrics_csv,
        qa_csv=qa_csv,
        methodology_path=methodology_path,
        snapshot_sources=list(snapshot_sources),
        enhanced_folder=bundle.enhanced_folder,
    )


def main() -> None:
    args = parse_args()
    result = build_current_snapshot_analytics(
        enhanced_root=Path(args.enhanced_root).expanduser().resolve(),
        explicit_enhanced_folder=Path(args.enhanced_folder).expanduser().resolve() if args.enhanced_folder else None,
        explicit_enhanced_workbook=Path(args.enhanced_workbook).expanduser().resolve() if args.enhanced_workbook else None,
        snapshot_root=Path(args.snapshot_root).expanduser().resolve(),
        explicit_snapshot=Path(args.snapshot).expanduser().resolve() if args.snapshot else None,
        output_root=Path(args.output_root).expanduser().resolve(),
    )
    print(f"Current snapshot analytics created in: {result.output_folder}")
    print(f"Workbook: {result.workbook_path}")
    print(f"Snapshot sources used: {len(result.snapshot_sources)}")


if __name__ == "__main__":
    main()
