from __future__ import annotations

import argparse
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Sequence, Tuple

from openpyxl import Workbook

from src.build_master_roster import STATUS_PRIORITY, autosize_columns, is_excluded_chapter, style_header
from src.canonical_bundle import DEFAULT_CANONICAL_ROOT, load_canonical_bundle


ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUTPUT_DIR = ROOT / "Yearly"


@dataclass(frozen=True)
class ChapterMember:
    chapter: str
    last_name: str
    first_name: str
    banner_id: str
    status: str

    def as_list(self) -> List[str]:
        return [self.last_name, self.first_name, self.banner_id, self.status]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Build yearly chapter roster workbooks from the canonical roster_term table. "
            "Each year gets its own workbook, with one sheet per chapter."
        )
    )
    parser.add_argument("--canonical-root", default=str(DEFAULT_CANONICAL_ROOT))
    parser.add_argument("--canonical-folder", default="")
    parser.add_argument("-o", "--output-dir", default=str(DEFAULT_OUTPUT_DIR))
    return parser.parse_args()


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def safe_sheet_name(value: str) -> str:
    cleaned = clean_text(value) or "Unknown"
    for char in "[]:*?/\\":  # Excel-invalid sheet chars
        cleaned = cleaned.replace(char, "")
    return cleaned[:31] or "Unknown"


def choose_preferred_member(existing: ChapterMember, candidate: ChapterMember) -> ChapterMember:
    existing_unknown = existing.chapter.lower() == "unknown"
    candidate_unknown = candidate.chapter.lower() == "unknown"
    if existing_unknown and not candidate_unknown:
        return candidate
    if STATUS_PRIORITY.get(candidate.status, 10) > STATUS_PRIORITY.get(existing.status, 10):
        return candidate
    if bool(candidate.banner_id) and not bool(existing.banner_id):
        return candidate
    return existing


def dedupe_chapter_members(rows: Iterable[ChapterMember]) -> List[ChapterMember]:
    best_rows: Dict[Tuple[str, str, str, str], ChapterMember] = {}
    for row in rows:
        if row.banner_id:
            key = (row.chapter.lower(), "banner", row.banner_id.lower(), "")
        else:
            key = (row.chapter.lower(), "name", row.last_name.lower(), row.first_name.lower())
        existing = best_rows.get(key)
        best_rows[key] = row if existing is None else choose_preferred_member(existing, row)
    return list(best_rows.values())


def rows_to_yearly_chapters(roster_term) -> Dict[str, Dict[str, List[ChapterMember]]]:
    grouped: Dict[str, Dict[str, List[ChapterMember]]] = defaultdict(lambda: defaultdict(list))
    for row in roster_term.itertuples(index=False):
        year_value = clean_text(getattr(row, "term_year", ""))
        chapter = clean_text(getattr(row, "chapter", "")) or "Unknown"
        if not year_value or year_value.lower() == "unknown":
            continue
        if is_excluded_chapter(chapter):
            continue
        last_name = clean_text(getattr(row, "last_name", ""))
        first_name = clean_text(getattr(row, "first_name", ""))
        banner_id = clean_text(getattr(row, "student_id", ""))
        status = clean_text(getattr(row, "org_status_bucket", ""))
        if not any([last_name, first_name, banner_id]):
            continue
        grouped[year_value][chapter].append(
            ChapterMember(
                chapter=chapter,
                last_name=last_name,
                first_name=first_name,
                banner_id=banner_id,
                status=status,
            )
        )
    return grouped


def write_year_workbook(year_label: str, chapter_rows: Dict[str, Sequence[ChapterMember]], output_path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    for chapter in sorted(chapter_rows):
        ws = wb.create_sheet(title=safe_sheet_name(chapter))
        ws.append(["Last Name", "First Name", "Banner ID", "Status"])
        style_header(ws)
        for member in sorted(
            dedupe_chapter_members(chapter_rows[chapter]),
            key=lambda item: (item.last_name.lower(), item.first_name.lower(), item.banner_id.lower()),
        ):
            ws.append(member.as_list())
        ws.freeze_panes = "A2"
        autosize_columns(ws)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def build_yearly_chapter_rosters(canonical_root: Path, explicit_folder: Path | None, output_dir: Path) -> None:
    bundle = load_canonical_bundle(canonical_root=canonical_root, explicit_folder=explicit_folder)
    roster_term = bundle.tables["roster_term"].copy()
    if roster_term.empty:
        raise FileNotFoundError("No usable canonical roster_term rows were found.")

    grouped = rows_to_yearly_chapters(roster_term)
    if not grouped:
        raise FileNotFoundError("No yearly chapter rows were found in the canonical roster_term table.")

    output_dir.mkdir(parents=True, exist_ok=True)
    for year_label in sorted(grouped, key=lambda value: int(value) if value.isdigit() else 9999):
        write_year_workbook(
            year_label=year_label,
            chapter_rows=grouped[year_label],
            output_path=output_dir / f"{year_label}.xlsx",
        )


def main() -> None:
    args = parse_args()
    build_yearly_chapter_rosters(
        canonical_root=Path(args.canonical_root).expanduser().resolve(),
        explicit_folder=Path(args.canonical_folder).expanduser().resolve() if args.canonical_folder else None,
        output_dir=Path(args.output_dir).expanduser().resolve(),
    )
    print(f"Yearly chapter rosters created in: {Path(args.output_dir).expanduser().resolve()}")


if __name__ == "__main__":
    main()
