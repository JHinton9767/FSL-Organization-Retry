from __future__ import annotations

import argparse
from collections import defaultdict
from pathlib import Path
from typing import Dict, List, Sequence, Tuple

from openpyxl import Workbook

from src.build_master_roster import autosize_columns, style_header
from src.canonical_bundle import DEFAULT_CANONICAL_ROOT, load_canonical_bundle


ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUTPUT_WORKBOOK = ROOT / "Master_Roster_Grades.xlsx"

COMBINED_COLUMNS = [
    "Academic Year",
    "Term",
    "Chapter",
    "Last Name",
    "First Name",
    "Banner ID",
    "Email",
    "Roster Present",
    "Academic Present",
    "Status",
    "Position",
    "New Member Flag",
    "Join Term",
    "Join Basis",
    "Major",
    "Attempted Hours",
    "Passed Hours",
    "Total Cumulative Hours",
    "Academic Standing",
    "Term GPA",
    "Institutional Cumulative GPA",
    "Overall Cumulative GPA",
    "Final Outcome",
    "Exit Reason",
    "Resolved Outcome Flag",
    "Outcome Evidence Source",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Build Master_Roster_Grades.xlsx as a downstream export from the canonical analytics bundle. "
            "This workbook is no longer an analytical source of truth."
        )
    )
    parser.add_argument("--canonical-root", default=str(DEFAULT_CANONICAL_ROOT))
    parser.add_argument("--canonical-folder", default="")
    parser.add_argument("-o", "--output", default=str(DEFAULT_OUTPUT_WORKBOOK))
    parser.add_argument("--chunk-size", type=int, default=1000)
    return parser.parse_args()


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def combined_rows(master_longitudinal) -> List[List[object]]:
    rows: List[List[object]] = []
    ordered = master_longitudinal.sort_values(
        by=["observed_year", "observed_term_sort", "chapter", "last_name", "first_name", "student_id"],
        ascending=[True, True, True, True, True, True],
        na_position="last",
    )
    for row in ordered.itertuples(index=False):
        rows.append(
            [
                clean_text(getattr(row, "observed_year", "")),
                clean_text(getattr(row, "term_label", "")),
                clean_text(getattr(row, "chapter", "")),
                clean_text(getattr(row, "last_name", "")),
                clean_text(getattr(row, "first_name", "")),
                clean_text(getattr(row, "student_id", "")),
                clean_text(getattr(row, "email", "")),
                clean_text(getattr(row, "roster_present", "")),
                clean_text(getattr(row, "academic_present", "")),
                clean_text(getattr(row, "org_status_bucket", "")),
                clean_text(getattr(row, "org_position_raw", "")),
                clean_text(getattr(row, "new_member_flag", "")),
                clean_text(getattr(row, "join_term", "")),
                clean_text(getattr(row, "org_entry_term_basis", "")),
                clean_text(getattr(row, "major", "")),
                getattr(row, "attempted_hours_term", ""),
                getattr(row, "earned_hours_term", ""),
                getattr(row, "total_cumulative_hours", ""),
                clean_text(getattr(row, "academic_standing_bucket", "")),
                getattr(row, "term_gpa", ""),
                getattr(row, "institutional_cumulative_gpa", ""),
                getattr(row, "overall_cumulative_gpa", ""),
                clean_text(getattr(row, "final_outcome_bucket", "")),
                clean_text(getattr(row, "exit_reason_code", "")),
                clean_text(getattr(row, "resolved_outcome_flag", "")),
                clean_text(getattr(row, "outcome_evidence_source", "")),
            ]
        )
    return rows


def write_summary_sheet(wb: Workbook, rows: Sequence[List[object]], source_folder: Path) -> None:
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Metric", "Value"])
    style_header(ws)

    matched_roster = sum(1 for row in rows if clean_text(row[7]) == "Yes")
    matched_academic = sum(1 for row in rows if clean_text(row[8]) == "Yes")
    resolved_outcomes = sum(1 for row in rows if clean_text(row[24]) == "Yes")
    metrics = [
        ["Canonical source folder", str(source_folder)],
        ["Combined student-term rows", len(rows)],
        ["Rows with roster data", matched_roster],
        ["Rows with academic data", matched_academic],
        ["Rows with resolved final outcomes", resolved_outcomes],
        ["Important note", "This workbook is a downstream export from canonical tables and is not the analytical source of truth."],
    ]
    for metric in metrics:
        ws.append(metric)
    ws.freeze_panes = "A2"
    autosize_columns(ws)


def write_combined_sheets(wb: Workbook, rows: Sequence[List[object]], chunk_size: int) -> None:
    grouped: Dict[Tuple[str, str], List[List[object]]] = defaultdict(list)
    for row in rows:
        grouped[(clean_text(row[0]), clean_text(row[1]))].append(row)

    for academic_year, term in sorted(grouped.keys(), key=lambda item: (int(item[0]) if str(item[0]).isdigit() else 9999, item[1])):
        semester_rows = grouped[(academic_year, term)]
        for start in range(0, len(semester_rows), chunk_size):
            end = min(start + chunk_size, len(semester_rows))
            chunk_number = (start // chunk_size) + 1
            sheet_name = f"{term} {chunk_number}"[:31]
            ws = wb.create_sheet(title=sheet_name)
            ws.append(COMBINED_COLUMNS)
            style_header(ws)
            for row in semester_rows[start:end]:
                ws.append(row)
            ws.freeze_panes = "A2"
            autosize_columns(ws)


def build_master_roster_grades(canonical_root: Path, explicit_folder: Path | None, output_path: Path, chunk_size: int) -> None:
    bundle = load_canonical_bundle(canonical_root=canonical_root, explicit_folder=explicit_folder)
    master_longitudinal = bundle.tables["master_longitudinal"].copy()
    rows = combined_rows(master_longitudinal)
    if not rows:
        raise FileNotFoundError("No canonical master_longitudinal rows were available for export.")

    wb = Workbook()
    write_summary_sheet(wb, rows, bundle.output_folder)
    write_combined_sheets(wb, rows, chunk_size=chunk_size)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main() -> None:
    args = parse_args()
    build_master_roster_grades(
        canonical_root=Path(args.canonical_root).expanduser().resolve(),
        explicit_folder=Path(args.canonical_folder).expanduser().resolve() if args.canonical_folder else None,
        output_path=Path(args.output).expanduser().resolve(),
        chunk_size=args.chunk_size,
    )
    print(f"Master roster/grades workbook created: {Path(args.output).expanduser().resolve()}")


if __name__ == "__main__":
    main()
